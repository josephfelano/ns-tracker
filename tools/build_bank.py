"""Build the NetSuite tracker question bank (bank/*.json) from the canonical workbooks.

Usage:
    python -I tools/build_bank.py --source "C:/path/to/workbooks" [--out bank] [--generated-at ISO]

--source  folder that contains "NS GUIDE/" and netsuite-readiness-packet-v1.xlsx (read-only).
--out     output folder (default: <repo>/bank).
--generated-at  override meta.generatedAt. Default is the newest source-file mtime (UTC), so
          re-running on unchanged sources produces byte-identical output.

Only bank (template) content is emitted. Answers, owners and statuses are instance data and
are deliberately excluded. Corrections are applied as explicit, asserted patches and logged
to bank/CORRECTIONS.md.
"""
import argparse
import datetime as dt
import hashlib
import json
import os
import re
import sys

import openpyxl

BANK_VERSION = "1.0.0"

SRC = {
    "discovery": "NS GUIDE/netsuite-discovery-workbook.xlsx",
    "register": "NS GUIDE/netsuite-decision-register.xlsx",
    "prework": "NS GUIDE/netsuite-implementation-readiness-prework.xlsx",
    "readinessV1": "netsuite-readiness-packet-v1.xlsx",
    "packetV2": "netsuite-preimplementation-packet-v2.md",
    "auditNotes": "netsuite-readiness-packet-audit-notes.md",
}

AUDIT = "netsuite-readiness-packet-audit-notes.md"
QID_RE = re.compile(r"\b([A-Z0-9]{3})-(\d+)([a-z]?)\b")


# --------------------------------------------------------------------------- helpers
def die(msg):
    sys.stderr.write("build_bank: ERROR: " + msg + "\n")
    sys.exit(1)


def s(v):
    """Normalise a cell value to a stripped string or None."""
    if v is None:
        return None
    v = str(v).strip()
    return v or None


def is_example(cell):
    rgb = cell.fill.fgColor.rgb if cell.fill and cell.fill.fgColor is not None else None
    return isinstance(rgb, str) and rgb.upper().endswith("F0F0F0")


def is_header(cell):
    rgb = cell.fill.fgColor.rgb if cell.fill and cell.fill.fgColor is not None else None
    return isinstance(rgb, str) and rgb.upper().endswith("1F3243")


def split_list_validation(formula1):
    """Split an Excel literal list. Excel treats every comma as a separator, but the
    workbooks contain one value with an embedded comma ("Available, cleanup needed").
    Tokens that start with a space are re-joined to the previous token; see CORRECTIONS C-04."""
    raw = formula1.strip()
    if raw.startswith('"') and raw.endswith('"'):
        raw = raw[1:-1]
    out = []
    for tok in raw.split(","):
        if out and tok.startswith(" "):
            out[-1] = out[-1] + "," + tok
        else:
            out.append(tok)
    return [t.strip() for t in out]


def validations(ws):
    """Map column letter(s) of each list validation -> values."""
    res = {}
    for dv in ws.data_validations.dataValidation:
        if dv.type != "list" or not dv.formula1:
            continue
        for rng in str(dv.sqref).split():
            col = re.match(r"([A-Z]+)", rng).group(1)
            res[col] = split_list_validation(dv.formula1)
    return res


def header_row(ws, row):
    return [s(c.value) for c in ws[row]]


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        h.update(f.read())
    return h.hexdigest()


def write_json(out_dir, name, data):
    path = os.path.join(out_dir, name)
    with open(path, "w", encoding="utf-8", newline="\n") as f:
        f.write(json.dumps(data, indent=2, ensure_ascii=False))
        f.write("\n")
    return path


# --------------------------------------------------------------------------- packet v2 (prefix titles)
def parse_packet_prefixes(path):
    parts, prefixes, cur_part = {}, [], None
    with open(path, encoding="utf-8") as f:
        for line in f:
            m = re.match(r"^# PART (\d) — (.+?)\s*$", line)
            if m:
                cur_part = "P" + m.group(1)
                continue
            m = re.match(r"^## \d+\. ([A-Z0-9]{3}) — (.+?)\s*$", line)
            if m:
                prefixes.append({"prefix": m.group(1), "title": m.group(2), "part": cur_part})
    if len(prefixes) != 13:
        die("expected 13 prefix headings in packet v2, got %d" % len(prefixes))
    return prefixes


# --------------------------------------------------------------------------- discovery
PART_SHEETS = [("P1", "Part 1 - Structural"), ("P2", "Part 2 - Process"), ("P3", "Part 3 - Wrap Up")]
EXPECTED_Q_HEADER = ["ID", "Tier", "Section", "Question", "Why we're asking", "Answer", "File / link", "Owner", "Status"]


def parse_discovery(path):
    wb = openpyxl.load_workbook(path)
    questions, sections, status_enum = [], [], None
    for part_id, sheet in PART_SHEETS:
        ws = wb[sheet]
        hdr = header_row(ws, 4)[: len(EXPECTED_Q_HEADER)]
        if hdr != EXPECTED_Q_HEADER:
            die("unexpected header on %s: %r" % (sheet, hdr))
        v = validations(ws)
        if status_enum is None:
            status_enum = v["I"]
        elif v["I"] != status_enum:
            die("status enum differs on " + sheet)
        for r in range(5, ws.max_row + 1):
            qid = s(ws.cell(r, 1).value)
            if not qid:
                continue
            m = QID_RE.fullmatch(qid)
            if not m:
                die("bad question id %r on %s row %d" % (qid, sheet, r))
            section = s(ws.cell(r, 3).value)
            if not any(x["name"] == section and x["part"] == part_id for x in sections):
                sections.append({"name": section, "part": part_id, "sheet": sheet, "prefixes": []})
            sec = next(x for x in sections if x["name"] == section and x["part"] == part_id)
            if m.group(1) not in sec["prefixes"]:
                sec["prefixes"].append(m.group(1))
            q = {
                "id": qid,
                "parentId": "%s-%s" % (m.group(1), m.group(2)) if m.group(3) else None,
                "prefix": m.group(1),
                "part": part_id,
                "section": section,
                "tier1": s(ws.cell(r, 2).value) == "Tier 1",
                "text": s(ws.cell(r, 4).value),
                "why": s(ws.cell(r, 5).value),
            }
            questions.append(q)

    # Systems tab tables and Reports tab
    ws = wb["Systems"]
    tables, cur = [], None
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1)
        av = s(a.value)
        m = re.match(r"^Table (\d) - (.+?)\s+\((.+)\)$", av or "")
        if m:
            cur = {"key": None, "number": int(m.group(1)), "title": m.group(2),
                   "questionIds": [x.strip() for x in m.group(3).replace(" and ", ",").split(",")],
                   "description": s(ws.cell(r + 1, 1).value), "columns": [], "example": None,
                   "presetRows": []}
            tables.append(cur)
            continue
        if cur is None or av is None:
            continue
        if is_header(a):
            cur["columns"] = [x for x in header_row(ws, r) if x]
        elif is_example(a):
            cur["example"] = {cur["columns"][i]: s(ws.cell(r, i + 1).value)
                              for i in range(len(cur["columns"])) if s(ws.cell(r, i + 1).value)}
        elif cur["columns"] and r > 0:
            cur["presetRows"].append(av)
    keys = {1: "systemInventory", 2: "systemOfRecord", 3: "criticalSpreadsheets", 4: "interfaces"}
    for t in tables:
        t["key"] = keys[t["number"]]
    ws = wb["Reports"]
    reports = {"key": "reports", "title": "Recurring reports", "questionIds": ["RPT-1"],
               "description": s(ws.cell(2, 1).value), "columns": [x for x in header_row(ws, 4) if x],
               "example": None}
    if is_example(ws.cell(5, 1)):
        reports["example"] = {reports["columns"][i]: s(ws.cell(5, i + 1).value)
                              for i in range(len(reports["columns"])) if s(ws.cell(5, i + 1).value)}

    # File checklist
    ws = wb["File Checklist"]
    if header_row(ws, 4)[:5] != ["Category", "File", "Notes", "Filename / link", "Status"]:
        die("unexpected File Checklist header")
    fc_status = validations(ws)["E"]
    files = []
    for r in range(5, ws.max_row + 1):
        cat, fil = s(ws.cell(r, 1).value), s(ws.cell(r, 2).value)
        if cat and fil:
            files.append({"category": cat, "file": fil, "notes": s(ws.cell(r, 3).value)})
    return questions, sections, status_enum, tables, reports, files, fc_status


# --------------------------------------------------------------------------- decision register
REG_HEADER = ["Question ID", "Tier", "Area", "What the answer decides", "NetSuite object / setting",
              "Reversibility", "Layer", "Blocks / notes", "Owner (theirs)", "Status", "Answer / decision"]


def parse_register(path):
    wb = openpyxl.load_workbook(path)
    ws = wb["Decision Register"]
    if header_row(ws, 4)[:11] != REG_HEADER:
        die("unexpected Decision Register header")
    dec_status = validations(ws)["J"]
    decisions = []
    for r in range(5, ws.max_row + 1):
        raw = s(ws.cell(r, 1).value)
        rev = s(ws.cell(r, 6).value)
        if not raw or not rev:  # progress block below the table has no reversibility
            continue
        decisions.append({
            "_raw": raw,
            "tier1": s(ws.cell(r, 2).value) == "Tier 1",
            "area": s(ws.cell(r, 3).value),
            "decides": s(ws.cell(r, 4).value),
            "nsSetting": s(ws.cell(r, 5).value),
            "reversibility": rev,
            "layer": s(ws.cell(r, 7).value),
            "notes": s(ws.cell(r, 8).value),
        })

    ws = wb["One-Way Doors"]
    if header_row(ws, 4)[:5] != ["#", "Constraint", "Why it matters here", "Confidence", "Source"]:
        die("unexpected One-Way Doors header")
    doors = []
    for r in range(5, ws.max_row + 1):
        if ws.cell(r, 1).value is None:
            continue
        c = s(ws.cell(r, 2).value)
        doors.append({"num": int(ws.cell(r, 1).value), "constraint": c,
                      "why": s(ws.cell(r, 3).value), "confidence": s(ws.cell(r, 4).value),
                      "source": s(ws.cell(r, 5).value),
                      "isOneWay": not c.upper().startswith("NOT A ONE-WAY DOOR")})

    ws = wb["Build Sequence"]
    if header_row(ws, 4)[:5] != ["Layer", "What gets decided", "Cost of changing it later",
                                 "Fed by packet sections", "Notes"]:
        die("unexpected Build Sequence header")
    layers = []
    for r in range(5, ws.max_row + 1):
        code = s(ws.cell(r, 1).value)
        if code:
            layers.append({"code": code, "decides": s(ws.cell(r, 2).value),
                           "costOfChange": s(ws.cell(r, 3).value),
                           "fedByRaw": s(ws.cell(r, 4).value), "notes": s(ws.cell(r, 5).value)})

    ws = wb["Module Scoping"]
    if header_row(ws, 4)[:5] != ["Ask", "Driven by", "Why it matters", "Answer / quoted?", "Status"]:
        die("unexpected Module Scoping header")
    ms_status = validations(ws)["E"]
    scoping = []
    for r in range(5, ws.max_row + 1):
        ask = s(ws.cell(r, 1).value)
        if ask:
            scoping.append({"ask": ask, "drivenByRaw": s(ws.cell(r, 2).value),
                            "why": s(ws.cell(r, 3).value)})
    return decisions, dec_status, doors, layers, scoping, ms_status


# --------------------------------------------------------------------------- prework
def parse_prework(path):
    wb = openpyxl.load_workbook(path)
    enums = {}
    ws = wb["Data & Documents"]
    dd_cols = [x for x in header_row(ws, 4) if x]
    v = validations(ws)
    enums["preworkReadiness"], enums["preworkCollectionStatus"] = v["F"], v["I"]
    data_docs = []
    for r in range(5, ws.max_row + 1):
        area, item = s(ws.cell(r, 1).value), s(ws.cell(r, 2).value)
        if area and item:
            data_docs.append({"area": area, "item": item, "goodEnough": s(ws.cell(r, 3).value)})

    ws = wb["Process Inventory"]
    pi_cols = [x for x in header_row(ws, 4) if x]
    v = validations(ws)
    enums["preworkDocumentation"], enums["preworkClarity"] = v["E"], v["F"]
    if v["I"] != enums["preworkCollectionStatus"]:
        die("collection status enum differs on Process Inventory")
    procs = []
    for r in range(5, ws.max_row + 1):
        area, proc = s(ws.cell(r, 1).value), s(ws.cell(r, 2).value)
        if area and proc:
            procs.append({"area": area, "process": proc})

    ws = wb["Systems & Partners"]
    sp_cols = [x for x in header_row(ws, 4) if x]
    v = validations(ws)
    enums["systemType"], enums["systemExchange"] = v["A"], v["F"]
    enums["systemDocsAvailable"], enums["systemSampleAvailable"] = v["G"], v["H"]
    enums["systemExpectedAfter"] = v["I"]
    if v["K"] != enums["preworkCollectionStatus"]:
        die("collection status enum differs on Systems & Partners")
    enum_by_col = {"A": "systemType", "F": "systemExchange", "G": "systemDocsAvailable",
                   "H": "systemSampleAvailable", "I": "systemExpectedAfter",
                   "K": "preworkCollectionStatus"}
    sys_template = []
    for i, name in enumerate(sp_cols):
        col = openpyxl.utils.get_column_letter(i + 1)
        sys_template.append({"column": name, "enum": enum_by_col.get(col)})
    dd_enum = {"Data readiness": "preworkReadiness", "Collection status": "preworkCollectionStatus"}
    pi_enum = {"Existing documentation": "preworkDocumentation", "Current-state clarity": "preworkClarity",
               "Collection status": "preworkCollectionStatus"}
    return (data_docs, [{"column": c, "enum": dd_enum.get(c)} for c in dd_cols],
            procs, [{"column": c, "enum": pi_enum.get(c)} for c in pi_cols], sys_template, enums)


def parse_readiness_v1(path):
    ws = openpyxl.load_workbook(path)["Readiness Checklist"]
    hdr = header_row(ws, 4)
    if hdr[0] != "ID" or hdr[8] != "Prepares (Discovery Workbook)":
        die("unexpected Readiness Checklist header")
    out = {}
    for r in range(5, ws.max_row + 1):
        rid = s(ws.cell(r, 1).value)
        if rid:
            raw = s(ws.cell(r, 9).value) or ""
            out[rid] = {"raw": raw, "ids": ["%s-%s%s" % m for m in QID_RE.findall(raw)],
                        "tier1Tagged": ["%s-%s%s" % m.groups()
                                        for m in re.finditer(QID_RE.pattern + r"\s*\[Tier 1\]", raw)]}
    return out


# --------------------------------------------------------------------------- crosswalks (heuristic, human-confirmable)
# Prework Data & Documents item -> readiness-packet-v1 rows whose "Prepares" refs it inherits.
# Optional regex narrows the inherited refs when one v1 row covered two prework items.
PREWORK_CROSSWALK = {
    "Legal entity / subsidiary list": (["R-FIN-1"], r"^FIN-1"),
    "Fiscal calendar": (["R-FIN-1"], r"^FIN-2"),
    "Chart of accounts": (["R-FIN-2"], r"^DIM-[12]$"),
    "Reporting dimensions": (["R-FIN-2"], r"^DIM-3$"),
    "Management reporting package": (["R-FIN-5"], None),
    "Month-end close checklist": (["R-FIN-4"], None),
    "Fixed asset register": (["R-FIN-7"], None),
    "Cash & reconciliation account list": (["R-SYS-5"], None),
    "Item / SKU master": (["R-ITM-1", "R-ITM-2"], None),
    "BOMs / recipes / formulas": (["R-BOM-1", "R-BOM-2"], None),
    "UOM / packaging / barcode data": (["R-ITM-3"], None),
    "Customer master": (["R-CUS-1"], None),
    "Vendor / partner master": (["R-VEN-1"], None),
    "Inventory location list": (["R-INV-1"], None),
    "Price lists / pricing files": (["R-CUS-3"], None),
    "Product / ingredient specifications": (["R-QLT-2"], None),
    "Employee / user roster": (["R-ORG-2"], None),
    "System / application inventory": (["R-SYS-1"], None),
    "System-of-record map": (["R-SYS-2"], None),
    "Interface / integration inventory": (["R-SYS-3"], None),
    "Critical spreadsheet inventory": (["R-SYS-4"], None),
    "Recurring report inventory": (["R-RPT-1", "R-RPT-2"], r"^RPT-"),
    "Approval / control documents": (["R-CTL-1"], None),
    "Quality / traceability documents": (["R-QLT-3"], None),
}

# File Checklist item (prefix of the File text) -> prework Data & Documents item, with coverage.
FILE_CROSSWALK = {
    "Chart of accounts export": ("Chart of accounts", "full"),
    "Last close checklist": ("Month-end close checklist", "partial"),
    "Monthly reporting package": ("Management reporting package", "full"),
    "Item / SKU export": ("Item / SKU master", "full"),
    "Packaging hierarchy": ("UOM / packaging / barcode data", "full"),
    "Customer export": ("Customer master", "full"),
    "Vendor export": ("Vendor / partner master", "full"),
    "Price lists": ("Price lists / pricing files", "full"),
    "Bills of material": ("BOMs / recipes / formulas", "full"),
    "Ingredient and packaging specifications": ("Product / ingredient specifications", "full"),
    "Approved supplier list": ("Product / ingredient specifications", "full"),
    "Employee list": ("Employee / user roster", "full"),
    "A/P platform vendor and bill export": ("Vendor / partner master", "partial"),
    "Most recent mock recall": ("Quality / traceability documents", "full"),
}

# One-way door # -> regexes matched against decision "decides | nsSetting".
DOOR_CROSSWALK = {
    1: [r"Costing method PER ITEM"],
    2: [r"Multi-Location Inventory"],
    3: [r"Multi-Location Inventory", r"Opening inventory quantity"],
    4: [r"Multi-Location Inventory"],
    5: [r"Enable Features > Lot Tracking", r"Lot/serial designation per item"],
    6: [r"Multiple Currencies", r"Advanced Revenue Management", r"Statistical Accounts",
        r"Advanced Taxes", r"Multi-Location Inventory"],
    7: [r"Enable Features > Multiple Units of Measure", r"Statistical Accounts"],
    8: [r"OneWorld and how many subsidiaries", r"provision OneWorld now"],
    9: [r"Enable Features > Advanced Receiving"],
    10: [r"Enable Features > Advanced Receiving", r"Receiving process incl\. quality hold"],
    11: [r"BOM structure, versioning"],
    12: [r"Which accounts exist and which legacy"],
    13: [r"sold by catch weight"],
    14: [r"Minimum remaining shelf life"],
    15: [r"Item type per SKU"],
    16: [r"Costing method PER ITEM", r"Item type per SKU"],
}


# --------------------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--source", required=True)
    ap.add_argument("--out", default=os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "bank"))
    ap.add_argument("--generated-at", default=None)
    a = ap.parse_args()

    paths = {k: os.path.join(a.source, v) for k, v in SRC.items()}
    for k, p in paths.items():
        if not os.path.isfile(p):
            die("missing source %s: %s" % (k, p))
    os.makedirs(a.out, exist_ok=True)

    prefixes = parse_packet_prefixes(paths["packetV2"])
    questions, sections, q_status, sys_tables, reports, files, fc_status = parse_discovery(paths["discovery"])
    decisions, d_status, doors, layers, scoping, ms_status = parse_register(paths["register"])
    data_docs, dd_template, procs, pi_template, sys_template, pw_enums = parse_prework(paths["prework"])
    v1 = parse_readiness_v1(paths["readinessV1"])

    qindex = {q["id"]: i for i, q in enumerate(questions)}
    if len(qindex) != len(questions):
        die("duplicate question ids")
    qbyid = {q["id"]: q for q in questions}
    order = lambda ids: sorted(set(ids), key=lambda x: qindex[x])  # noqa: E731

    corrections, normalizations, human = [], [], []

    # ---- C-01 Landed Cost decision keyed to FIN-4a -> FIN-4c
    hits = [d for d in decisions if d["_raw"] == "FIN-4a" and "Landed Cost" in (d["nsSetting"] or "")]
    if len(hits) != 1:
        die("C-01: expected exactly one Landed Cost decision keyed FIN-4a, found %d" % len(hits))
    hits[0]["_raw"] = "FIN-4c"
    corrections.append({"id": "C-01", "where": "decisions.json (Decision Register row 17, 'Whether freight/duty/handling capitalize into inventory')",
                        "what": "questionIds FIN-4a -> FIN-4c",
                        "why": "FIN-4a asks how inventory is valued; the capitalization question is FIN-4c. Module Scoping already keys the same decision to FIN-4c.",
                        "source": AUDIT + " lines 160-164 (section 7, item 1)"})

    # ---- C-02 FIN-4c Tier 1
    if qbyid["FIN-4c"]["tier1"]:
        die("C-02: FIN-4c already Tier 1; correction no longer needed")
    qbyid["FIN-4c"]["tier1"] = True
    corrections.append({"id": "C-02", "where": "questions.json FIN-4c (Discovery 'Part 1 - Structural' row 28)",
                        "what": "tier1 false -> true (Tier 1 rows 31 -> 32)",
                        "why": "The Landed Cost decision it drives is Tier 1 and Module Scoping lists Landed Cost as a pre-quote question.",
                        "source": AUDIT + " lines 166-169 (section 7, item 2)"})

    # ---- C-04 embedded-comma enum value
    if "Available, cleanup needed" not in pw_enums["preworkReadiness"]:
        die("C-04: expected 'Available, cleanup needed' in prework readiness enum")
    corrections.append({"id": "C-04", "where": "meta.json enums.preworkReadiness (Prework 'Data & Documents' F5:F28 dropdown)",
                        "what": "Kept 'Available, cleanup needed' as ONE value. In the workbook the dropdown splits it into 'Available' and ' cleanup needed' because Excel treats the comma as a separator.",
                        "why": "The sheet's conditional formatting and the Start Here COUNTIF both test for the full string 'Available, cleanup needed', so that is the intended value; in Excel it can never be selected.",
                        "source": "Found by inspecting the workbook's data validation (not in the audit notes)"})

    # ---- decisions: split multi-IDs, pseudo refs
    out_decisions = []
    for i, d in enumerate(decisions, 1):
        parts = [p.strip() for p in d["_raw"].split("/")]
        qids = [p for p in parts if p in qindex]
        pseudo = [p for p in parts if p not in qindex]
        for p in pseudo:
            if QID_RE.fullmatch(p):
                die("decision ref %r looks like a question id but does not exist" % p)
        out_decisions.append({
            "id": "D-%03d" % i, "questionIds": order(qids), "pseudoRefs": pseudo, "tier1": d["tier1"],
            "area": d["area"], "decides": d["decides"], "nsSetting": d["nsSetting"],
            "reversibility": d["reversibility"], "layer": d["layer"], "notes": d["notes"],
        })
    normalizations.append("Decision Register 'Question ID' values containing ' / ' were split into questionIds; non-question tokens (ART, CLS) are kept in pseudoRefs.")

    # ---- layers: expand ranges / parents
    def expand(token):
        token = token.strip()
        if token in qindex:
            return [token], None
        m = re.fullmatch(r"([A-Z0-9]{3})-(\d+) to \1-(\d+)", token)
        if m:
            ids = []
            for n in range(int(m.group(2)), int(m.group(3)) + 1):
                sub, _ = expand("%s-%d" % (m.group(1), n))
                ids += sub or []
            return ids, None
        kids = [q["id"] for q in questions if q["parentId"] == token]
        if kids:
            return kids, None
        return None, token

    out_layers = []
    parent_expanded = set()
    for L in layers:
        ids, pseudo = [], []
        for tok in L["fedByRaw"].split(","):
            got, bad = expand(tok)
            if got is None:
                if QID_RE.fullmatch(bad.strip()):
                    die("layer %s: unresolvable ref %r" % (L["code"], bad))
                pseudo.append(bad)
            else:
                ids += got
                if tok.strip() not in qindex and " to " not in tok:
                    parent_expanded.add(tok.strip())
        out_layers.append({"code": L["code"], "decides": L["decides"], "costOfChange": L["costOfChange"],
                           "fedByQuestionIds": order(ids), "fedByPseudoRefs": pseudo,
                           "fedByRaw": L["fedByRaw"], "notes": L["notes"]})
    corrections.append({"id": "C-03", "where": "layers.json fedByQuestionIds (Build Sequence 'Fed by packet sections')",
                        "what": "Parent IDs that are not workbook rows (%s) and ranges ('X-1 to X-n') expanded to the real child question IDs. The original text is kept in fedByRaw." % ", ".join(sorted(parent_expanded, key=lambda t: (t.split('-')[0], int(t.split('-')[1])))),
                        "why": "Build Sequence cites parent IDs that do not exist as workbook rows, which breaks ID-based lookup.",
                        "source": AUDIT + " lines 171-173 (section 7, item 3)"})

    # ---- one-way doors -> decisions (heuristic)
    out_doors = []
    for dr in doors:
        pats = DOOR_CROSSWALK.get(dr["num"])
        if pats is None:
            die("no door crosswalk for #%d" % dr["num"])
        ids = []
        for p in pats:
            m = [d["id"] for d in out_decisions if re.search(p, "%s | %s" % (d["decides"], d["nsSetting"]))]
            if not m:
                die("door #%d pattern %r matched no decision" % (dr["num"], p))
            ids += m
        dr = dict(dr)
        dr["decisionIds"] = sorted(set(ids))
        dr["mappedBy"] = "heuristic"
        out_doors.append(dr)

    # ---- module scoping
    out_scoping = []
    for i, m in enumerate(scoping, 1):
        raw = m["drivenByRaw"] or ""
        ids = [t.strip() for t in raw.split(",") if t.strip() in qindex]
        rest = [t.strip() for t in raw.split(",") if t.strip() and t.strip() not in qindex]
        for t in rest:
            if QID_RE.fullmatch(t):
                die("module scoping ref %r does not exist" % t)
        out_scoping.append({"id": "M-%02d" % i, "ask": m["ask"], "drivenByQuestionIds": order(ids),
                            "drivenByAllTier1": raw == "All Tier 1", "drivenByRaw": raw, "why": m["why"]})

    # ---- prework
    out_dd = []
    for i, d in enumerate(data_docs, 1):
        if d["item"] not in PREWORK_CROSSWALK:
            die("no prework crosswalk for %r" % d["item"])
        rows, flt = PREWORK_CROSSWALK[d["item"]]
        ids = []
        for rid in rows:
            if rid not in v1:
                die("readiness v1 row %s missing" % rid)
            ids += [x for x in v1[rid]["ids"] if not flt or re.search(flt, x)]
        for x in ids:
            if x not in qindex:
                die("prework prepares ref %s missing" % x)
        out_dd.append({"id": "PW-D%02d" % i, "area": d["area"], "item": d["item"], "goodEnough": d["goodEnough"],
                       "preparesQuestionIds": order(ids), "preparesFrom": rows, "mappedBy": "heuristic"})
    dd_by_item = {d["item"]: d["id"] for d in out_dd}
    out_procs = [{"id": "PW-P%02d" % i, "area": p["area"], "process": p["process"]} for i, p in enumerate(procs, 1)]

    # ---- file checklist
    out_files, used = [], set()
    for i, f in enumerate(files, 1):
        rec = {"id": "F-%02d" % i, "category": f["category"], "file": f["file"], "notes": f["notes"]}
        for pre, (item, cov) in FILE_CROSSWALK.items():
            if f["file"].startswith(pre):
                rec["coveredByPreworkId"] = dd_by_item[item]
                rec["coverage"] = cov
                rec["mappedBy"] = "heuristic"
                used.add(pre)
        out_files.append(rec)
    if used != set(FILE_CROSSWALK):
        die("file crosswalk entries unused: %r" % (set(FILE_CROSSWALK) - used))

    # ---- strip SYS-2 example duplicate / trailing "Other" row from presets
    for t in sys_tables:
        if t["key"] == "systemOfRecord":
            presets = [p for p in t["presetRows"] if not p.lower().startswith("other")]
            t["allowOtherRows"] = len(presets) != len(t["presetRows"])
            t["presetDataTypes"] = presets
        else:
            if t["presetRows"]:
                die("unexpected preset rows in Systems table %s" % t["key"])
        del t["presetRows"]

    # ---- consistency findings for humans (computed, not guessed)
    t1q = {q["id"] for q in questions if q["tier1"]}
    for d in out_decisions:
        if d["tier1"] and d["questionIds"] and not (set(d["questionIds"]) & t1q):
            human.append("%s ('%s') is Tier 1 but none of its questions (%s) are Tier 1." % (d["id"], d["decides"], ", ".join(d["questionIds"])))
    t1_dec_q = {x for d in out_decisions if d["tier1"] for x in d["questionIds"]}
    for qid in sorted(t1q - t1_dec_q, key=lambda x: qindex[x]):
        human.append("Question %s is Tier 1 but no Tier 1 decision references it (it may only sharpen another Tier 1 answer)." % qid)
    for d in out_decisions:
        if not d["tier1"] and set(d["questionIds"]) & t1q and d["layer"] == "L0":
            human.append("%s ('%s', L0) is not Tier 1 although it references Tier 1 question(s) %s." % (d["id"], d["decides"], ", ".join(sorted(set(d["questionIds"]) & t1q, key=lambda x: qindex[x]))))
    for m in out_scoping:
        non = [x for x in m["drivenByQuestionIds"] if x not in t1q]
        if non:
            human.append("%s ('%s') is a pre-signature ask but is driven by non-Tier-1 question(s) %s. Same reasoning as audit item 2 (FIN-4c), but the audit flagged only FIN-4c." % (m["id"], m["ask"], ", ".join(non)))
    for rid, info in v1.items():
        for x in info["tier1Tagged"]:
            if x in qbyid and not qbyid[x]["tier1"]:
                human.append("readiness-packet-v1 %s tags %s as [Tier 1] but discovery does not." % (rid, x))

    # ---- meta
    mt = max(os.path.getmtime(p) for p in paths.values())
    gen = a.generated_at or dt.datetime.fromtimestamp(mt, dt.timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")
    part_names = {pid: sh for pid, sh in PART_SHEETS}
    meta = {
        "version": BANK_VERSION,
        "generatedAt": gen,
        "generatedAtBasis": "newest source file mtime (override with --generated-at)",
        "sources": [{"key": k, "file": SRC[k].replace("\\", "/"), "sha256": sha256(paths[k])} for k in SRC],
        "enums": {
            "questionStatus": q_status,
            "decisionStatus": d_status,
            "reversibility": ["One-way", "Costly", "Easy"],
            "layer": [L["code"] for L in out_layers],
            "decisionArea": sorted({d["area"] for d in out_decisions}),
            "preworkReadiness": pw_enums["preworkReadiness"],
            "preworkCollectionStatus": pw_enums["preworkCollectionStatus"],
            "preworkDocumentation": pw_enums["preworkDocumentation"],
            "preworkClarity": pw_enums["preworkClarity"],
            "preworkDataArea": list(dict.fromkeys(d["area"] for d in out_dd)),
            "preworkProcessArea": list(dict.fromkeys(p["area"] for p in out_procs)),
            "systemType": pw_enums["systemType"],
            "systemExchange": pw_enums["systemExchange"],
            "systemDocsAvailable": pw_enums["systemDocsAvailable"],
            "systemSampleAvailable": pw_enums["systemSampleAvailable"],
            "systemExpectedAfter": pw_enums["systemExpectedAfter"],
            "moduleScopingStatus": ms_status,
            "fileChecklistStatus": fc_status,
            "fileChecklistCategory": list(dict.fromkeys(f["category"] for f in out_files)),
            "oneWayDoorConfidence": list(dict.fromkeys(d["confidence"] for d in out_doors)),
        },
        "statusSemantics": {
            "questionDone": ["Complete", "Not applicable"],
            "questionOutstandingVisible": ["Need help"],
            "decisionDone": ["Decided", "N/A"],
        },
        "parts": [{"id": pid, "name": part_names[pid],
                   "packetTitle": {"P1": "Structural", "P2": "Process", "P3": "Validation"}[pid]}
                  for pid, _ in PART_SHEETS],
        "prefixes": prefixes,
        "sections": sections,
    }
    if q_status != meta["enums"]["fileChecklistStatus"]:
        die("file checklist status differs from question status")

    write_json(a.out, "meta.json", meta)
    write_json(a.out, "questions.json", questions)
    write_json(a.out, "decisions.json", out_decisions)
    write_json(a.out, "layers.json", out_layers)
    write_json(a.out, "oneWayDoors.json", out_doors)
    write_json(a.out, "moduleScoping.json", out_scoping)
    write_json(a.out, "prework.json", {"dataDocuments": out_dd, "dataDocumentsTemplate": dd_template,
                                       "processes": out_procs, "processesTemplate": pi_template,
                                       "systemsTemplate": sys_template})
    write_json(a.out, "systemsTables.json", {"tables": sys_tables, "reports": reports})
    write_json(a.out, "fileChecklist.json", out_files)

    # ---- CORRECTIONS.md
    static_human = [
        "Org chart missing from the File Checklist (" + AUDIT + " lines 175-176, section 7 item 4). ORG-1 asks for it, but adding a 38th checklist row is a content change, so it was NOT added.",
        "ITM-1a on the 'Costing method PER ITEM' decision (" + AUDIT + " lines 178-180, section 7 item 5). The decision stays Tier 1 because FIN-4a is Tier 1; ITM-1a stays non-Tier-1 and stays linked. Confirm ITM-1a ('attach the item export') should remain a questionId at all.",
        "Parent-ID expansion in layers.json is over-inclusive by design: e.g. L0 'FIN-4' pulls in FIN-4a..FIN-4d and 'ITM-6' pulls in ITM-6a..ITM-6h, although only some of those children drive L0 settings. Narrow by hand if layer membership needs to be exact.",
        "Pseudo refs left unresolved: decisions use 'ART' and 'ART / CLS' (cutover rows); Build Sequence L7 uses 'ART, CLS'. Decide whether ART means ART-1 + the File Checklist, and whether bare CLS means the whole Close section (CLS-1..CLS-11).",
        "Module Scoping M-01 ('All Tier 1') is flagged drivenByAllTier1 rather than listing IDs; M-18 (renewal uplift cap) is driven by '-' (no question).",
        "All crosswalks marked mappedBy:'heuristic' need a human pass: oneWayDoors.decisionIds, prework.dataDocuments.preparesQuestionIds (inherited from readiness-packet-v1 'Prepares' via a hand crosswalk in tools/build_bank.py), and fileChecklist.coveredByPreworkId (14 of 37 covered; README says 'roughly fifteen').",
        "The Discovery 'Start Here' text says 'about nineteen' Tier 1 rows (that is the question count; there are now 32 rows after C-02). Questions were not reworded per instructions.",
    ]
    lines = ["# Bank corrections log", "",
             "Generated by `tools/build_bank.py`. Do not edit by hand; edit the builder and re-run.", "",
             "Sources: `NS GUIDE/` discovery workbook, decision register, implementation-readiness-prework; "
             "`netsuite-readiness-packet-v1.xlsx` (Prepares mapping only); `netsuite-preimplementation-packet-v2.md` (prefix titles).", "",
             "## Corrections applied", ""]
    for c in sorted(corrections, key=lambda c: c["id"]):
        lines += ["### %s" % c["id"], "", "- **What:** " + c["what"], "- **Where:** " + c["where"],
                  "- **Why:** " + c["why"], "- **Source:** " + c["source"], ""]
    lines += ["## Normalizations (structure only, no content change)", ""]
    lines += ["- " + n for n in normalizations + [
        "Grey example rows (fill F0F0F0) are emitted only as `example` objects in systemsTables.json, never as data rows.",
        "SYS-2 has 13 preset data types plus an 'Other (add rows as needed)' row, emitted as allowOtherRows:true. The grey example row (a duplicate of 'Inventory on hand - by location') is the table's `example`.",
        "Decision IDs D-001..D-087, scoping IDs M-01..M-18, prework IDs PW-D01..PW-D24 / PW-P01..PW-P23 and file IDs F-01..F-37 are synthetic and assigned in workbook row order.",
    ]]
    lines += ["", "## Needs a human decision", ""]
    lines += ["- " + h for h in static_human]
    lines += ["", "### Computed Tier 1 consistency findings", ""]
    lines += ["- " + h for h in human] if human else ["- none"]
    with open(os.path.join(a.out, "CORRECTIONS.md"), "w", encoding="utf-8", newline="\n") as f:
        f.write("\n".join(lines) + "\n")

    print("questions=%d tier1=%d decisions=%d layers=%d doors=%d scoping=%d dataDocs=%d processes=%d files=%d humanFindings=%d" % (
        len(questions), sum(q["tier1"] for q in questions), len(out_decisions), len(out_layers), len(out_doors),
        len(out_scoping), len(out_dd), len(out_procs), len(out_files), len(human) + len(static_human)))


if __name__ == "__main__":
    main()
